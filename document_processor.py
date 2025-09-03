import asyncio
import re
from typing import List, Dict, Any
from docx import Document
import openai
from config import config, prompts

class DocumentProcessor:
    def __init__(self):
        self.client = openai.AsyncOpenAI(
            api_key=config.OPENAI_API_KEY,
            base_url=config.OPENAI_BASE_URL
        )
    
    def extract_text_from_docx(self, file_path: str) -> str:
        """从docx文件中提取文本内容"""
        try:
            doc = Document(file_path)
            text_content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text.strip())
            
            # 处理表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_content.append(" | ".join(row_text))
            
            return "\n".join(text_content)
        except Exception as e:
            raise Exception(f"文档读取失败: {str(e)}")

    def extract_text_from_pdf(self, file_path: str) -> str:
        """从PDF文件中提取文本内容"""
        try:
            # 延迟导入以避免在未安装依赖时阻塞应用启动
            import pdfplumber  # type: ignore
            texts: List[str] = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        texts.append(page_text.strip())
            return "\n".join(texts)
        except ImportError:
            raise Exception("缺少pdfplumber依赖，无法解析PDF。请安装pdfplumber或移除该文件。")
        except Exception as e:
            raise Exception(f"PDF读取失败: {str(e)}")

    def extract_text_from_excel(self, file_path: str) -> str:
        """从Excel文件（.xlsx/.xls）中提取文本内容"""
        try:
            if file_path.lower().endswith('.xlsx'):
                import openpyxl  # type: ignore
                wb = openpyxl.load_workbook(file_path, data_only=True)
                texts: List[str] = []
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    texts.append(f"[工作表] {sheet_name}")
                    for row in ws.iter_rows(values_only=True):
                        row_values = [str(c).strip() for c in row if c is not None and str(c).strip()]
                        if row_values:
                            texts.append(" | ".join(row_values))
                return "\n".join(texts)
            else:
                import xlrd  # type: ignore
                book = xlrd.open_workbook(file_path)
                texts: List[str] = []
                for sheet in book.sheets():
                    texts.append(f"[工作表] {sheet.name}")
                    for r in range(sheet.nrows):
                        row_values = [str(v).strip() for v in sheet.row_values(r) if str(v).strip()]
                        if row_values:
                            texts.append(" | ".join(row_values))
                return "\n".join(texts)
        except ImportError:
            raise Exception("缺少openpyxl/xlrd依赖，无法解析Excel。请安装openpyxl和xlrd或移除该文件。")
        except Exception as e:
            raise Exception(f"Excel读取失败: {str(e)}")

    def extract_text_from_doc(self, file_path: str) -> str:
        """从旧版Word .doc 文件中提取文本内容（使用textract，可能依赖系统工具）"""
        try:
            import textract  # type: ignore
            raw: bytes = textract.process(file_path)
            text = raw.decode('utf-8', errors='ignore').strip()
            return text
        except ImportError:
            raise Exception("缺少textract依赖，无法解析.doc。请安装textract或将文件转换为.docx。")
        except Exception as e:
            raise Exception(f".doc读取失败: {str(e)}")

    def extract_text_from_any(self, file_path: str) -> str:
        """根据文件后缀提取文本内容，支持 .pdf .docx .doc .xlsx .xls"""
        lower = file_path.lower()
        if lower.endswith('.pdf'):
            return self.extract_text_from_pdf(file_path)
        if lower.endswith('.docx'):
            return self.extract_text_from_docx(file_path)
        if lower.endswith('.doc'):
            return self.extract_text_from_doc(file_path)
        if lower.endswith('.xlsx') or lower.endswith('.xls'):
            return self.extract_text_from_excel(file_path)
        raise Exception("不支持的文件类型，仅支持.pdf .docx .doc .xlsx .xls")
    
    def chunk_text(self, text: str) -> List[str]:
        """将文本分块，优化分块策略"""
        if not text.strip():
            return []
        
        # 按段落分割
        paragraphs = [p.strip() for p in text.split('\n') if p.strip()]
        
        chunks = []
        current_chunk = ""
        
        for paragraph in paragraphs:
            # 如果当前段落加上现有块超过最大长度
            if len(current_chunk) + len(paragraph) + 1 > config.MAX_CHUNK_SIZE:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                    current_chunk = ""
                
                # 如果单个段落就超过最大长度，需要进一步分割
                if len(paragraph) > config.MAX_CHUNK_SIZE:
                    # 按句子分割
                    sentences = re.split(r'[。！？.!?]', paragraph)
                    temp_chunk = ""
                    
                    for sentence in sentences:
                        if sentence.strip():
                            sentence = sentence.strip() + "。"
                            if len(temp_chunk) + len(sentence) > config.MAX_CHUNK_SIZE:
                                if temp_chunk:
                                    chunks.append(temp_chunk.strip())
                                temp_chunk = sentence
                            else:
                                temp_chunk += sentence
                    
                    if temp_chunk.strip():
                        current_chunk = temp_chunk
                else:
                    current_chunk = paragraph
            else:
                if current_chunk:
                    current_chunk += "\n" + paragraph
                else:
                    current_chunk = paragraph
        
        if current_chunk.strip():
            chunks.append(current_chunk.strip())
        
        return chunks
    
    async def summarize_chunk(self, chunk: str, chunk_index: int) -> str:
        """对单个文本块进行总结"""
        try:
            prompt = prompts.SUMMARIZE_CHUNK.format(chunk=chunk)
            
            response = await self.client.chat.completions.create(
                model=config.OPENAI_MODEL,
                messages=[
                    {"role": "system", "content": "你是一名专业的投标专员，任务是精准地从标书文件中提取和总结关键信息。"},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=8192,
                temperature=0.3
            )
            
            return response.choices[0].message.content.strip()
        
        except Exception as e:
            return f"第{chunk_index + 1}块总结失败: {str(e)}"
    
    async def process_chunks_parallel(self, chunks: List[str]) -> List[str]:
        """并行处理所有文本块的总结"""
        semaphore = asyncio.Semaphore(config.MAX_CONCURRENT_REQUESTS)
        
        async def process_single_chunk(chunk: str, index: int) -> str:
            async with semaphore:
                return await self.summarize_chunk(chunk, index)
        
        tasks = [process_single_chunk(chunk, i) for i, chunk in enumerate(chunks)]
        summaries = await asyncio.gather(*tasks, return_exceptions=True)
        
        # 处理异常结果
        processed_summaries = []
        for i, summary in enumerate(summaries):
            if isinstance(summary, Exception):
                processed_summaries.append(f"第{i + 1}块处理失败: {str(summary)}")
            else:
                processed_summaries.append(summary)
        
        return processed_summaries
    
    async def generate_final_summary(self, chunk_summaries: List[str]) -> str:
        """生成最终的综合总结"""
        try:
            combined_summaries = "\n\n".join([f"### 第 {i+1} 部分总结\n{summary}" for i, summary in enumerate(chunk_summaries)])
            
            prompt = prompts.GENERATE_FINAL_SUMMARY.format(combined_summaries=combined_summaries)
            
            response = await self.client.chat.completions.create(
                model=config.OPENAI_MODEL,
                messages=[
                    {"role": "system", "content": "你是一名资深的投标策略师，擅长将零散信息整合成结构化的分析报告。"},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=8192,
                temperature=0.3
            )
            
            return response.choices[0].message.content.strip()
        
        except Exception as e:
            raise Exception(f"生成最终总结失败: {str(e)}")
    
    async def process_document(self, file_path: str) -> Dict[str, Any]:
        """完整的文档处理流程"""
        try:
            # 1. 提取文本
            text_content = self.extract_text_from_docx(file_path)
            if not text_content.strip():
                raise Exception("文档内容为空")
            
            # 2. 文本分块
            chunks = self.chunk_text(text_content)
            if not chunks:
                raise Exception("文档分块失败")
            
            # 3. 并行处理各块总结
            chunk_summaries = await self.process_chunks_parallel(chunks)
            
            # 4. 生成最终总结
            final_summary = await self.generate_final_summary(chunk_summaries)
            
            return {
                "success": True,
                "original_text_length": len(text_content),
                "chunks_count": len(chunks),
                "chunk_summaries": chunk_summaries,
                "final_summary": final_summary,
                "processing_info": {
                    "max_chunk_size": config.MAX_CHUNK_SIZE,
                    "chunks_processed": len(chunks),
                    "parallel_processing": True
                }
            }
        
        except Exception as e:
            return {
                "success": False,
                "error": str(e),
                "final_summary": None
            } 