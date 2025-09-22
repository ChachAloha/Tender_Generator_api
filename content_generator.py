import asyncio
import os
from typing import Dict, List, Any, Optional
import openai
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from config import config, prompts

class ContentGenerator:
    def __init__(self, model_config: Optional[Dict[str, str]] = None):
        api_key = (model_config or {}).get("api_key", config.OPENAI_API_KEY)
        base_url = (model_config or {}).get("base_url", config.OPENAI_BASE_URL)
        self.model_name = (model_config or {}).get("model", config.OPENAI_MODEL)
        self.client = openai.AsyncOpenAI(
            api_key=api_key,
            base_url=base_url
        )
    
    def _flatten_outline_to_sections(self, outline_data: Dict) -> List[Dict]:
        """将嵌套的目录结构扁平化为所有层级的章节列表"""
        sections = []
        outline_items = outline_data.get("outline", [])
        
        def extract_sections(items: List[Dict], parent_path: str = "", parent_id: Optional[str] = None):
            for item in items:
                item_id = item.get("id", "")
                # 创建当前章节的路径
                current_path = f"{parent_path} > {item.get('title', '')}" if parent_path else item.get('title', '')
                
                children = item.get("children", [])

                # 添加当前章节
                section = {
                    "id": item_id,
                    "title": item.get("title", ""),
                    "level": item.get("level", 1),
                    "parent_path": parent_path,
                    "full_path": current_path,
                    "parent_id": parent_id,
                    "has_children": bool(children)
                }
                sections.append(section)
                
                # 递归处理子章节
                if children:
                    extract_sections(children, current_path, item_id)
        
        extract_sections(outline_items)
        return sections
    
    async def generate_section_content(self, section: Dict, summary: str) -> Dict:
        """为单个章节生成内容，分两次生成并合并，实现续写效果。"""
        try:
            section_title = section.get("title", "未命名章节")
            section_level = section.get("level", 1)
            parent_path = section.get("parent_path", "")
            full_path = section.get("full_path", section_title)
            parent_id = section.get("parent_id")
            has_children = section.get("has_children", False)

            # Level 1 章节不生成内容
            # Level 2 章节如果有子章节，也不生成内容
            if section_level == 1 or (section_level == 2 and has_children):
                return {
                    "id": section.get("id", ""),
                    "title": section_title,
                    "level": section_level,
                    "content": "",
                    "parent_path": parent_path,
                    "full_path": full_path,
                    "parent_id": parent_id,
                    "has_children": has_children
                }

            # 构建上下文信息
            context_info = ""
            if parent_path:
                context_info = f"章节路径：{full_path}\n父章节：{parent_path}\n"
            else:
                context_info = f"章节路径：{full_path}\n"
            
            # --- 第一次生成 ---
            prompt1 = prompts.GENERATE_SECTION_CONTENT.format(
                section_title=section_title,
                section_level=section_level,
                context_info=context_info,
                summary=summary
            )
            
            response1 = await self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": "你是一个专业的内容创作者。请严格按照用户要求的格式输出内容，使用<p></p>标记段落，不要使用其他任何格式。"},
                    {"role": "user", "content": prompt1}
                ],
                max_tokens=8192,
                temperature=0.3
            )
            
            content1_raw = response1.choices[0].message.content or ""
            content1_normalized = self._normalize_content(content1_raw.strip())

            # 如果第一次生成的内容为空，则直接返回，不进行续写
            if not content1_normalized:
                return {
                    "id": section.get("id", ""),
                    "title": section_title,
                    "level": section_level,
                    "content": "", # 返回空内容
                    "parent_path": parent_path,
                    "full_path": full_path,
                    "parent_id": parent_id,
                    "has_children": has_children
                }

            # --- 第二次生成 (续写) ---
            prompt2 = prompts.CONTINUE_SECTION_CONTENT.format(
                section_title=section_title,
                section_level=section_level,
                context_info=context_info,
                summary=summary,
                existing_content=content1_normalized
            )

            response2 = await self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": "你是一个专业的内容创作者，擅长在已有内容的基础上进行扩写和续写。请严格按照用户要求的格式输出内容，使用<p></p>标记段落，不要使用其他任何格式。"},
                    {"role": "user", "content": prompt2}
                ],
                max_tokens=8192,
                temperature=0.4
            )

            content2_raw = response2.choices[0].message.content or ""
            content2_normalized = self._normalize_content(content2_raw.strip())

            # --- 合并内容 ---
            final_content = content1_normalized
            if content2_normalized:
                final_content += "\n\n" + content2_normalized
            
            return {
                "id": section.get("id", ""),
                "title": section_title,
                "level": section_level,
                "content": final_content,
                "parent_path": parent_path,
                "full_path": full_path,
                "parent_id": parent_id,
                "has_children": has_children
            }
        
        except Exception as e:
            return {
                "id": section.get("id", ""),
                "title": section_title,
                "level": section_level,
                "content": f"内容生成失败: {str(e)}",
                "parent_path": parent_path,
                "full_path": full_path,
                "parent_id": section.get("parent_id"),
                "has_children": section.get("has_children", False)
            }
    
    async def generate_all_sections_parallel(self, outline_data: Dict, summary: str, section_generator_func=None) -> List[Dict]:
        """并行生成所有层级章节的内容"""
        # 1. 将嵌套结构扁平化为所有章节的列表
        all_sections = self._flatten_outline_to_sections(outline_data)
        
        if not all_sections:
            return []
        
        # 确定使用哪个章节生成函数
        generator_func = section_generator_func if section_generator_func else self.generate_section_content

        # 2. 使用信号量控制并发
        semaphore = asyncio.Semaphore(config.MAX_CONCURRENT_REQUESTS)
        
        async def process_section_with_semaphore(section: Dict) -> Dict:
            async with semaphore:
                return await generator_func(section, summary)
        
        # 3. 创建所有章节的任务列表
        tasks = [process_section_with_semaphore(section) for section in all_sections]
        
        # 4. 并行执行所有任务
        results = await asyncio.gather(*tasks, return_exceptions=True)
        
        # 5. 处理可能的异常
        processed_results = []
        for i, result in enumerate(results):
            if isinstance(result, Exception):
                # 如果发生异常，创建错误信息
                section = all_sections[i]
                processed_results.append({
                    "id": section.get("id", f"section_{i+1}"),
                    "title": section.get("title", f"章节 {i+1}"),
                    "level": section.get("level", 1),
                    "content": f"内容生成失败: {str(result)}",
                    "parent_path": section.get("parent_path", ""),
                    "full_path": section.get("full_path", section.get("title", f"章节 {i+1}")),
                    "parent_id": section.get("parent_id"),
                    "has_children": section.get("has_children", False)
                })
            else:
                processed_results.append(result)
        
        return processed_results
    
    def _rebuild_nested_structure(self, flat_sections: List[Dict]) -> List[Dict]:
        """将扁平化的章节列表重新构建为嵌套结构，用于文档生成"""
        section_map = {s["id"]: s for s in flat_sections}
        for s in section_map.values():
            s["children"] = []

        nested_structure = []
        for section in flat_sections:
            parent_id = section.get("parent_id")
            if parent_id and parent_id in section_map:
                section_map[parent_id]["children"].append(section)
            else:
                nested_structure.append(section)
        
        return nested_structure

    def _count_sections(self, sections: List[Dict]) -> int:
        """计算章节总数（扁平化计算）"""
        return len(sections)

    def _to_chinese_numeral(self, n: int) -> str:
        """将整数转换为中文数字。"""
        mapping = {
            1: '一', 2: '二', 3: '三', 4: '四', 5: '五', 6: '六', 7: '七', 8: '八', 9: '九', 10: '十',
            11: '十一', 12: '十二', 13: '十三', 14: '十四', 15: '十五', 16: '十六', 17: '十七', 18: '十八', 19: '十九', 20: '二十'
        }
        return mapping.get(n, str(n))

    def _get_heading_prefix(self, level: int, counters: List[int], style_template: str) -> str:
        """根据样式模板获取标题前缀。"""
        if level <= 1:
            return ""
        
        prefix = ""
        # counters 的长度由调用方 _add_sections_recursively 保证等于 level
        if style_template == 'A':
            if level == 3: 
                prefix = f"{counters[2]}、"
        elif style_template == 'B':
            if level == 3: 
                prefix = f"{self._to_chinese_numeral(counters[2])}、"
            elif level == 4: 
                prefix = f"{counters[3]}、"
        elif style_template == 'C':
            if level == 3: 
                prefix = f"（{counters[2]}）"
            elif level == 4: 
                prefix = f"{counters[3]}."
        elif style_template == 'D':
            if level == 3: 
                prefix = f"{counters[2]}."
            elif level == 4: 
                prefix = f"{counters[2]}.{counters[3]}"
        elif style_template == 'E':
            if level == 3: 
                prefix = f"（{self._to_chinese_numeral(counters[2])}）"
            elif level == 4: 
                prefix = f"（{counters[3]}）"
        
        return prefix

    def _apply_paragraph_style(self, p, run, level: int, style_template: str):
        """应用段落和字体样式。"""
        p.paragraph_format.line_spacing = 1.5

        def set_run_font(run, size_pt, bold=False, font_name='宋体'):
            run.font.name = font_name
            r = run._element.rPr.rFonts
            r.set(qn('w:eastAsia'), font_name)
            run.font.size = Pt(size_pt)
            run.font.bold = bold

        # 正文
        if level == 0:
            set_run_font(run, 12)
            # 设置首行缩进（2个字符）
            p.paragraph_format.first_line_indent = Pt(24)
            return

        # 总标题 - 所有样式模板都是二号居中加黑宋体
        if level == -1:  # 用于文档总标题
            set_run_font(run, 22, bold=True)  # 二号字
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return

        # 章节大标题 (Level 1) - 所有样式模板都是小二加黑宋体居中
        if level == 1:
            set_run_font(run, 18, bold=True)  # 小二号字
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return

        # 一级标题 (Level 2)
        if level == 2:
            if style_template in ['A', 'B']:
                set_run_font(run, 15, bold=True)  # 小三号字
                if style_template == 'B':
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # style_template A: 顶格（默认）
            elif style_template in ['C', 'D', 'E']:
                set_run_font(run, 12, bold=True)  # 小四号字
                if style_template == 'D':
                    p.paragraph_format.left_indent = Pt(24)  # 退格
                # style_template C 和 E: 顶格（默认）
            return
        
        # 其他级别标题 (Level 3-5)
        if level >= 3:
            set_run_font(run, 12)  # 小四号字
            # 根据样式模板可能需要调整缩进
            if style_template == 'D' and level == 4:
                p.paragraph_format.left_indent = Pt(24)  # 退格

    def _add_sections_recursively(self, doc: Document, sections: List[Dict], style_template: str, counters: List[int], parent_level: int = 0):
        """递归添加章节到文档。"""
        for i, section in enumerate(sections):
            current_level = section.get("level", 1)
            title = section.get("title", "")
            content = section.get("content", "")
            
            # 根据当前层级构建正确的counters
            # 确保counters的长度与current_level匹配
            if current_level > len(counters):
                # 如果当前层级大于counters长度，扩展counters
                while len(counters) < current_level:
                    counters.append(0)
            
            # 更新当前层级的计数
            if current_level <= len(counters):
                # 创建新的counters副本，避免修改原始列表
                new_counters = counters[:current_level]
                if current_level > 0:
                    new_counters[current_level - 1] = i + 1
            else:
                new_counters = counters + [i + 1]

            prefix = self._get_heading_prefix(current_level, new_counters, style_template)
            heading_text = f"{prefix}{title}"
            
            p = doc.add_paragraph()
            run = p.add_run(heading_text)
            self._apply_paragraph_style(p, run, current_level, style_template)

            # 使用新的内容添加方法
            self._add_content_to_doc(doc, content, style_template)

            children = section.get("children", [])
            if children:
                # 为子章节递归调用，传递正确的counters
                self._add_sections_recursively(doc, children, style_template, new_counters, current_level)
    
    def create_docx_from_template(self, sections: List[Dict], outline_data: Dict, output_path: str, style_template: str = 'A') -> str:
        """使用模板创建Word文档"""
        try:
            doc = Document()
            
            # 设置默认字体
            style = doc.styles['Normal']
            font = style.font
            font.name = '宋体'
            style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

            # 总标题
            title_p = doc.add_paragraph()
            title_run = title_p.add_run(outline_data.get("title", "生成的文档"))
            self._apply_paragraph_style(title_p, title_run, -1, style_template)
            
            # 重新构建嵌套结构用于文档生成
            nested_sections = self._rebuild_nested_structure(sections)
            
            # 递归添加章节
            self._add_sections_recursively(doc, nested_sections, style_template, [])
            
            # 保存文档
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            doc.save(output_path)
            
            return output_path
        
        except Exception as e:
            import traceback
            tb_str = traceback.format_exc()
            raise Exception(f"文档生成失败: {str(e)}\n{tb_str}")
    
    async def generate_document(self, summary: str, outline_data: Dict, style_template: str = 'A') -> Dict[str, Any]:
        """完整的文档生成流程"""
        try:
            # 1. 并行生成所有层级章节内容（扁平化处理）
            sections = await self.generate_all_sections_parallel(outline_data, summary)
            
            # 2. 生成输出文件路径
            timestamp = int(asyncio.get_event_loop().time())
            output_filename = f"generated_document_{timestamp}.docx"
            output_path = os.path.join(config.OUTPUT_DIR, output_filename)
            
            # 3. 创建文档
            final_path = self.create_docx_from_template(sections, outline_data, output_path, style_template)
            
            return {
                "success": True,
                "document_path": final_path,
                "sections_count": self._count_sections(sections),
                "outline_title": outline_data.get("title", "生成的文档"),
                "sections": sections
            }
        
        except Exception as e:
            return {
                "success": False,
                "error": str(e),
                "document_path": None
            }

    async def generate_supplementary_document(self, user_request: str, style_template: str = 'A') -> Dict[str, Any]:
        """
        生成补充文档的完整流程。
        1. 根据用户请求动态生成一个简单的outline。
        2. 基于此outline生成约2000字的内容。
        3. 返回纯文本和Word文档。
        """
        try:
            # 1. 动态生成补充文档的Outline
            supplementary_outline_prompt = prompts.GENERATE_SUPPLEMENTARY_OUTLINE.format(
                user_request=user_request
            )
            
            response = await self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": "你是一个专业的文档架构师，严格按照JSON格式输出。"},
                    {"role": "user", "content": supplementary_outline_prompt}
                ],
                max_tokens=8192,
                temperature=0.1,
                response_format={"type": "json_object"}
            )
            
            import json
            raw_content = response.choices[0].message.content
            if not raw_content:
                raise Exception("补充文档目录生成失败：模型返回空内容。")

            try:
                outline_data = json.loads(raw_content)
            except json.JSONDecodeError as e:
                raise Exception(f"补充文档目录生成失败：无法解析JSON。错误: {e}")

            # 2. 并行生成所有章节内容
            # 注意：这里的context从summary改为summary + user_request，让内容生成更贴近用户需求
            generation_context = f"项目总结：\n{summary}\n\n用户补充需求：\n{user_request}"
            sections = await self.generate_all_sections_parallel(
                outline_data, 
                generation_context,
                section_generator_func=self.generate_supplementary_section_content
            )
            
            # 3. 生成输出文件路径
            timestamp = int(asyncio.get_event_loop().time())
            output_filename = f"supplementary_document_{timestamp}.docx"
            output_path = os.path.join(config.OUTPUT_DIR, output_filename)
            
            # 4. 创建Word文档
            final_path = self.create_docx_from_template(sections, outline_data, output_path, style_template)
            
            # 5. 生成纯文本内容
            plain_text_content = self._generate_plain_text(sections, outline_data)

            return {
                "success": True,
                "document_path": final_path,
                "plain_text": plain_text_content,
                "sections_count": self._count_sections(sections),
                "outline_title": outline_data.get("title", "生成的补充文档"),
                "sections": sections
            }

        except Exception as e:
            return {
                "success": False,
                "error": str(e),
                "document_path": None,
                "plain_text": None
            }

    async def generate_supplementary_section_content(self, section: Dict, context: str) -> Dict:
        """为补充文档的单个章节生成精简内容。"""
        try:
            section_title = section.get("title", "未命名章节")
            full_path = section.get("full_path", section_title)
            
            # 对level 1, 2, 3 的标题不生成内容，只对最底层的level 4生成
            if section.get("level", 0) < 4:
                return {
                    "id": section.get("id", ""), "title": section_title, "level": section.get("level", 1),
                    "content": "", "parent_path": section.get("parent_path", ""), "full_path": full_path,
                    "parent_id": section.get("parent_id"), "has_children": section.get("has_children", False)
                }

            prompt = prompts.GENERATE_SUPPLEMENTARY_SECTION_CONTENT.format(
                context=context,
                section_title=section_title,
                full_path=full_path
            )
            response = await self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": "你是一个专业的内容创作者。请严格按照用户要求的格式输出内容，使用<p></p>标记段落。"},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=8192,
                temperature=0.3
            )
            
            content_raw = response.choices[0].message.content or ""
            normalized_content = self._normalize_content(content_raw.strip())
            
            return {
                "id": section.get("id", ""), "title": section_title, "level": section.get("level", 1),
                "content": normalized_content, "parent_path": section.get("parent_path", ""),
                "full_path": full_path, "parent_id": section.get("parent_id"),
                "has_children": section.get("has_children", False)
            }
        
        except Exception as e:
            return {
                "id": section.get("id", ""), "title": section.get("title", "未命名章节"), "level": section.get("level", 1),
                "content": f"内容生成失败: {str(e)}", "parent_path": section.get("parent_path", ""),
                "full_path": section.get("full_path", ""), "parent_id": section.get("parent_id"),
                "has_children": section.get("has_children", False)
            }

    def _generate_plain_text(self, sections: List[Dict], outline_data: Dict) -> str:
        """从生成的章节数据中创建纯文本文档。"""
        
        def clean_html(raw_html):
            import re
            cleanr = re.compile('<.*?>')
            cleantext = re.sub(cleanr, '', raw_html)
            return cleantext

        text_content = []
        text_content.append(f"# {outline_data.get('title', '补充文档')}\n\n")

        nested_sections = self._rebuild_nested_structure(sections)

        def append_section_text(sections_list, level=1):
            for section in sections_list:
                title = section.get('title', '')
                content = section.get('content', '')
                
                # 添加标题，使用Markdown风格
                text_content.append(f"{'#' * section.get('level', 1)} {title}\n")
                
                # 添加内容
                if content:
                    # 清理内容中的HTML标签（如<p>）和多余的换行
                    cleaned_content = clean_html(content).replace('\n\n', '\n').strip()
                    text_content.append(f"{cleaned_content}\n\n")

                if "children" in section and section["children"]:
                    append_section_text(section["children"], level + 1)
        
        append_section_text(nested_sections)
        
        return "".join(text_content)

    def _normalize_content(self, content: str) -> str:
        """规范化AI生成的内容，确保适合Word文档"""
        import re
        
        # 1. 提取<p>标签中的内容
        paragraphs = re.findall(r'<p>(.*?)</p>', content, re.DOTALL)
        
        # 如果没有找到<p>标签，尝试按换行符分割
        if not paragraphs:
            # 移除多余的空白行
            lines = [line.strip() for line in content.split('\n') if line.strip()]
            # 将连续的非空行合并为段落
            paragraphs = []
            current_paragraph = []
            
            for line in lines:
                # 如果是短句或看起来像是列表项，合并到当前段落
                if len(line) < 50 or line[0].isdigit() or line[0] in '•·-':
                    current_paragraph.append(line)
                else:
                    # 如果当前段落有内容，保存它
                    if current_paragraph:
                        paragraphs.append(' '.join(current_paragraph))
                        current_paragraph = []
                    # 开始新段落
                    current_paragraph.append(line)
            
            # 保存最后一个段落
            if current_paragraph:
                paragraphs.append(' '.join(current_paragraph))
        
        # 2. 清理每个段落
        cleaned_paragraphs = []
        for para in paragraphs:
            # 移除段落内的换行符
            para = para.replace('\n', ' ').replace('\r', ' ')
            # 移除多余的空格
            para = re.sub(r'\s+', ' ', para)
            # 移除首尾空白
            para = para.strip()
            
            # 移除Markdown格式（如果AI仍然输出了）
            # 移除加粗标记
            para = re.sub(r'\*\*(.*?)\*\*', r'\1', para)
            para = re.sub(r'__(.*?)__', r'\1', para)
            # 移除斜体标记
            para = re.sub(r'\*(.*?)\*', r'\1', para)
            para = re.sub(r'_(.*?)_', r'\1', para)
            # 移除代码标记
            para = re.sub(r'`(.*?)`', r'\1', para)
            # 移除标题标记
            para = re.sub(r'^#+\s+', '', para)
            # 移除列表标记
            para = re.sub(r'^[-*+]\s+', '', para)
            para = re.sub(r'^\d+\.\s+', '', para)
            # 移除引用标记
            para = re.sub(r'^>\s+', '', para)
            # 移除链接
            para = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', para)
            # 移除图片
            para = re.sub(r'!\[([^\]]*)\]\([^\)]+\)', '', para)
            # 移除HTML标签（除了我们自己的<p>标签）
            para = re.sub(r'<(?!/?p\b)[^>]+>', '', para)
            # 移除表格分隔符
            para = re.sub(r'\|', ' ', para)
            # 移除多余的标点符号
            para = re.sub(r'[#*_`~]', '', para)
            
            # 再次清理多余空格
            para = re.sub(r'\s+', ' ', para).strip()
            
            # 只保留有实质内容的段落（至少20个字符）
            if len(para) >= 20:
                cleaned_paragraphs.append(para)
        
        # 3. 合并过短的段落
        final_paragraphs = []
        i = 0
        while i < len(cleaned_paragraphs):
            current = cleaned_paragraphs[i]
            
            # 如果当前段落太短（少于100字符）且不是最后一个段落
            if len(current) < 100 and i < len(cleaned_paragraphs) - 1:
                # 与下一个段落合并
                next_para = cleaned_paragraphs[i + 1] if i + 1 < len(cleaned_paragraphs) else ""
                if next_para:
                    current = current + " " + next_para
                    i += 1  # 跳过下一个段落
            
            final_paragraphs.append(current)
            i += 1
        
        # 4. 确保每个段落都是完整的句子
        final_cleaned = []
        for para in final_paragraphs:
            # 确保段落以句号、问号或感叹号结尾
            if para and not para[-1] in '.。!！?？':
                para += '。'
            final_cleaned.append(para)
        
        # 5. 使用特殊分隔符连接段落，便于后续处理
        return "\n\n".join(final_cleaned)
    
    def _add_content_to_doc(self, doc: Document, content: str, style_template: str):
        """将规范化的内容添加到文档中"""
        # 按双换行符分割段落
        paragraphs = content.split("\n\n")
        
        for para_text in paragraphs:
            if para_text.strip():
                # 添加段落
                p = doc.add_paragraph()
                run = p.add_run(para_text.strip())
                # 应用正文样式
                self._apply_paragraph_style(p, run, 0, style_template) 